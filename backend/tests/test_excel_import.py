"""Tests for the Excel import / export / template feature."""

from datetime import date
from io import BytesIO

import openpyxl

from tests.conftest import get_admin_token

EXCEL_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_book(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import_file(client, headers, content, filename="import.xlsx"):
    return client.post(
        "/api/v1/import/excel",
        headers=headers,
        files={"file": (filename, content, EXCEL_CT)},
    )


def _load_book(content):
    return openpyxl.load_workbook(BytesIO(content), data_only=True)


def _laptop_numbers(client, headers):
    resp = client.get("/api/v1/laptops/", headers=headers)
    assert resp.status_code == 200
    return [l["laptop_number"] for l in resp.json()]


# ---------------------------------------------------------------------------
# Auth / validation guards
# ---------------------------------------------------------------------------
def test_import_requires_auth(client):
    resp = client.post(
        "/api/v1/import/excel",
        files={"file": ("x.xlsx", b"", EXCEL_CT)},
    )
    assert resp.status_code == 401


def test_export_requires_auth(client):
    assert client.get("/api/v1/export/excel").status_code == 401


def test_template_requires_auth(client):
    assert client.get("/api/v1/export/template").status_code == 401


def test_invalid_extension_rejected(client, admin_headers):
    resp = client.post(
        "/api/v1/import/excel",
        headers=admin_headers,
        files={"file": ("data.csv", b"some,data", "text/csv")},
    )
    assert resp.status_code == 400
    assert "xlsx" in resp.json()["detail"].lower()


def test_empty_file_rejected(client, admin_headers):
    resp = _import_file(client, admin_headers, b"")
    assert resp.status_code == 422


def test_missing_required_columns_rejected(client, admin_headers):
    data = _make_book(["Foo", "Bar"], [["a", "b"]])
    resp = _import_file(client, admin_headers, data)
    assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Downloadable template / export
# ---------------------------------------------------------------------------
def test_template_download_contains_headers(client, admin_headers):
    resp = client.get("/api/v1/export/template", headers=admin_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(EXCEL_CT)
    assert "template" in resp.headers["content-disposition"]

    wb = _load_book(resp.content)
    headers = [c.value for c in wb["Import Template"][1]]
    assert headers[0] == "Laptop ID"
    assert "Laptop Name" in headers
    assert "Serial Number" in headers
    assert "Allocation Date" in headers
    assert "Domain" in headers
    assert len(headers) == 14


def test_export_returns_workbook(client, admin_headers):
    resp = client.get("/api/v1/export/excel", headers=admin_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(EXCEL_CT)
    assert "export" in resp.headers["content-disposition"]

    wb = _load_book(resp.content)
    assert "Laptop Allocation Report" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Instructions" in wb.sheetnames
    assert wb["Laptop Allocation Report"].max_row >= 21  # header + 20 seeded laptops


# ---------------------------------------------------------------------------
# Import behaviour
# ---------------------------------------------------------------------------
def test_import_creates_laptops_interns_allocations(client, admin_headers):
    headers = [
        "Asset", "Brand", "Serial Number", "Status",
        "Name", "Email", "Domain", "Time Shift",
        "Allocation Date", "Return Date",
    ]
    data = _make_book(
        headers,
        [
            ["ASSET-X1", "HP", "SN-X1", "FREE", "Alice Chen", "alice@x.com",
             "Data Analytics", "9:00 AM to 6:00 PM IST", date(2026, 9, 1), date(2026, 9, 30)],
            ["ASSET-X2", "Dell", "SN-X2", "Hand Over", "Bob Ray", "bob@x.com",
             "Full Stack", "2:00 PM to 6:00 PM IST", date(2026, 9, 1), None],
        ],
    )
    resp = _import_file(client, admin_headers, data)
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["inserted"] == 2
    assert body["errors"] == []

    nums = _laptop_numbers(client, admin_headers)
    assert "ASSET-X1" in nums
    assert "ASSET-X2" in nums

    laptops = client.get("/api/v1/laptops/", headers=admin_headers).json()
    lp_x1 = next(l for l in laptops if l["laptop_number"] == "ASSET-X1")
    assert lp_x1["brand"] == "HP"
    assert lp_x1["serial_number"] == "SN-X1"
    assert lp_x1["status"] == "FREE"

    interns = client.get("/api/v1/interns/", headers=admin_headers).json()
    names = [i["name"] for i in interns]
    assert "Alice Chen" in names
    assert "Bob Ray" in names
    alice = next(i for i in interns if i["name"] == "Alice Chen")
    assert alice["batch"] == "Morning"       # from 9 AM Time Shift
    bob = next(i for i in interns if i["name"] == "Bob Ray")
    assert bob["batch"] == "Afternoon"       # from 2 PM Time Shift


def test_import_upsert_is_idempotent_and_updates(client, admin_headers):
    headers = ["Asset", "Brand", "RAM", "Status"]
    row = ["ASSET-U1", "Lenovo", "8 GB", "Free"]
    first = _make_book(headers, [row])
    r1 = _import_file(client, admin_headers, first)
    assert r1.json()["inserted"] == 1
    assert r1.json()["updated"] == 0

    # Same asset again but with a changed RAM -> update in place, no duplicate
    second = _make_book(headers, [["ASSET-U1", "Lenovo", "16 GB", "Free"]])
    r2 = _import_file(client, admin_headers, second)
    body = r2.json()
    assert body["inserted"] == 0
    assert body["updated"] == 1
    assert body["skipped"] == 0

    nums = _laptop_numbers(client, admin_headers)
    assert nums.count("ASSET-U1") == 1

    laptops = client.get("/api/v1/laptops/", headers=admin_headers).json()
    lp = next(l for l in laptops if l["laptop_number"] == "ASSET-U1")
    assert lp["ram"] == "16 GB"


def test_bad_row_is_skipped_and_reported(client, admin_headers):
    headers = [
        "Asset", "Status", "Name", "Domain",
        "Allocation Date", "Return Date",
    ]
    data = _make_book(
        headers,
        [
            ["ASSET-R1", "Free", "Row One Guy", "Cloud", date(2026, 9, 1), None],
            ["ASSET-R2", "Free", "Bad Row Guy", "Cloud",
             date(2026, 9, 30), date(2026, 9, 1)],  # return before allocation
            ["ASSET-R3", "Free", "Row Three Guy", "Cloud", date(2026, 9, 2), None],
        ],
    )
    resp = _import_file(client, admin_headers, data)
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["total_rows"] == 3
    assert body["inserted"] == 2
    assert body["skipped"] == 1
    assert len(body["errors"]) == 1
    assert body["errors"][0]["row"] == 3
    assert "Return Date" in body["errors"][0]["error"]

    # The bad row's laptop must NOT have been created
    nums = _laptop_numbers(client, admin_headers)
    assert "ASSET-R2" not in nums
    assert "ASSET-R1" in nums
    assert "ASSET-R3" in nums


def test_import_records_returned_allocation(client, admin_headers):
    stats_before = client.get("/api/v1/dashboard/stats", headers=admin_headers).json()
    assert stats_before["returned_allocations"] == 0

    headers = ["Asset", "Status", "Name", "Domain", "Allocation Date", "Return Date"]
    data = _make_book(
        headers,
        [["ASSET-M1", "Free", "Mia Stone", "UI/UX", date(2026, 8, 3), date(2026, 8, 20)]],
    )
    resp = _import_file(client, admin_headers, data)
    assert resp.status_code == 200
    assert resp.json()["inserted"] == 1

    stats = client.get("/api/v1/dashboard/stats", headers=admin_headers).json()
    assert stats["returned_allocations"] == 1